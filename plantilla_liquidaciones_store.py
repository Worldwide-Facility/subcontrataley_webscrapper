"""Planilla original de Liquidaciones de Sueldo, una por periodo.

El portal deja de entregarla después de cargar las liquidaciones. Se guarda
el Excel tal como se descarga (con RUT y faena, sin rellenar) y las otras
ejecuciones de ese mes lo reutilizan.
"""

from __future__ import annotations

import logging
import os
import re
import tempfile
import unicodedata
from collections.abc import Callable
from pathlib import Path

from config import GCS_BUCKET_NAME

logger = logging.getLogger(__name__)

PREFIJO_GCS = "Subcontrataley/plantillas_liquidaciones"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ENV_FORZAR = "PLANTILLA_LIQ_FORZAR"


def _forzar_sobrescritura() -> bool:
    return os.environ.get(ENV_FORZAR, "").strip().lower() in ("1", "true", "si", "sí", "yes")


def _contar_filas(path: Path) -> int | None:
    """Filas con datos de la primera hoja (None si no se puede leer)."""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            return sum(
                1 for fila in ws.iter_rows(values_only=True)
                if any(v not in (None, "") for v in fila)
            )
        finally:
            wb.close()
    except Exception:
        return None


def _filas_guardadas(blob) -> int | None:
    try:
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "guardada.xlsx"
            blob.download_to_filename(str(dest))
            return _contar_filas(dest)
    except Exception:
        return None


def clave_periodo(periodo: str) -> str:
    """'septiembre 2026' → 'septiembre_2026'."""
    s = unicodedata.normalize("NFKD", str(periodo or "").strip().lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_")
    if not s:
        raise ValueError("Periodo vacío para la planilla de liquidaciones")
    return s


def ruta_gcs_plantilla_liquidaciones(periodo: str) -> str:
    return f"{PREFIJO_GCS}/{clave_periodo(periodo)}.xlsx"


def _bucket(bucket, bucket_name: str | None):
    if bucket is not None:
        return bucket
    from gcs_docs import _cliente

    nombre = bucket_name or GCS_BUCKET_NAME
    return _cliente().bucket(nombre)


def guardar_plantilla_liquidaciones(
    origen: Path,
    periodo: str,
    *,
    bucket=None,
    bucket_name: str | None = None,
) -> str:
    """
    Sube el Excel original. Hay que llamarla antes de rellenar o borrar filas.

    Solo se conserva la PRIMERA planilla guardada del mes: si ya existe una
    para el periodo no se sobrescribe (salvo PLANTILLA_LIQ_FORZAR=1).
    Devuelve la ruta GCS en ambos casos.
    """
    ruta = ruta_gcs_plantilla_liquidaciones(periodo)
    blob = _bucket(bucket, bucket_name).blob(ruta)
    forzar = _forzar_sobrescritura()
    if blob.exists():
        filas_nueva = _contar_filas(Path(origen))
        filas_guardada = _filas_guardadas(blob)
        if not forzar:
            logger.info(
                "Ya existe planilla de liquidaciones para %s (gs://.../%s); se conserva "
                "la primera del mes y NO se sobrescribe (filas guardada=%s, nueva=%s). "
                "Use %s=1 para forzar.",
                periodo,
                ruta,
                filas_guardada,
                filas_nueva,
                ENV_FORZAR,
            )
            return ruta
        logger.warning(
            "%s activo: se sobrescribe la planilla guardada de %s "
            "(filas guardada=%s, nueva=%s)",
            ENV_FORZAR,
            periodo,
            filas_guardada,
            filas_nueva,
        )
    blob.upload_from_filename(str(origen), content_type=_XLSX)
    logger.info(
        "Planilla de liquidaciones guardada para %s → gs://%s/%s",
        periodo,
        getattr(blob, "bucket", None).name if getattr(blob, "bucket", None) else (bucket_name or GCS_BUCKET_NAME),
        ruta,
    )
    return ruta


def cargar_plantilla_liquidaciones(
    periodo: str,
    dest_dir: Path,
    *,
    bucket=None,
    bucket_name: str | None = None,
) -> Path | None:
    """Copia local de la planilla guardada, o None si ese periodo aún no está."""
    ruta = ruta_gcs_plantilla_liquidaciones(periodo)
    blob = _bucket(bucket, bucket_name).blob(ruta)
    if not blob.exists():
        return None
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"liquidaciones_{clave_periodo(periodo)}.xlsx"
    blob.download_to_filename(str(dest))
    logger.info("Planilla de liquidaciones de %s leída desde gs://%s", periodo, ruta)
    return dest


def asegurar_plantilla_liquidaciones(
    periodo: str,
    dest_dir: Path,
    descargar: Callable[[], Path],
    *,
    bucket=None,
    bucket_name: str | None = None,
) -> Path:
    """
    Planilla del periodo para armar RUT por faena.

    Si ya se guardó al descargar liquidaciones, no vuelve al portal.
    Si no existe, descarga, la guarda y la devuelve.
    """
    guardada = cargar_plantilla_liquidaciones(
        periodo, dest_dir, bucket=bucket, bucket_name=bucket_name
    )
    if guardada is not None:
        logger.info(
            "Reutilizando planilla de liquidaciones guardada para %s (%s)",
            periodo,
            guardada.name,
        )
        return guardada

    logger.info("Sin planilla guardada para %s; se descarga del portal", periodo)
    try:
        bajada = Path(descargar())
    except Exception as exc:
        raise RuntimeError(
            f"No hay planilla de liquidaciones guardada para '{periodo}' "
            "y el portal no permitió descargarla. "
            "Ejecute primero la carga de liquidaciones de ese mes."
        ) from exc
    guardar_plantilla_liquidaciones(
        bajada, periodo, bucket=bucket, bucket_name=bucket_name
    )
    return bajada
